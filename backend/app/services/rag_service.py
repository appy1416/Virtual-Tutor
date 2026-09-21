import os
import re
from datetime import datetime, timezone
import chromadb
from sentence_transformers import SentenceTransformer
from app.db.mongodb import get_database
from bson import ObjectId
from app.services.llm_service import LLMService
from typing import List, Dict, Any, Optional

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None

try:
    import pptx
except ImportError:
    pptx = None

# Persistent ChromaDB location
CHROMA_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "chroma_db"))
os.makedirs(CHROMA_DIR, exist_ok=True)
chroma_client = chromadb.PersistentClient(path=CHROMA_DIR)
_embedding_model = None

def get_embedding_model():
    global _embedding_model
    if _embedding_model is None:
        try:
            _embedding_model = SentenceTransformer("all-MiniLM-L6-v2")
        except Exception as e:
            print(f"Warning loading SentenceTransformer model: {e}")
            _embedding_model = None
    return _embedding_model

def chunk_text(text: str, chunk_size: int = 500, overlap: int = 100) -> List[str]:
    chunks = []
    if not text:
        return chunks
    start = 0
    text_len = len(text)
    while start < text_len:
        end = start + chunk_size
        chunks.append(text[start:end])
        start += (chunk_size - overlap)
        if start >= text_len or (chunk_size - overlap) <= 0:
            break
    return chunks

def extract_text_from_file(file_path: str, material_type: str) -> List[Dict[str, Any]]:
    pages = []
    if not os.path.exists(file_path):
        return pages

    if material_type == "text" or file_path.endswith(".txt"):
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
                pages.append({"page": 1, "text": content})
        except Exception as e:
            print(f"Error reading text file: {e}")
            
    elif material_type == "pdf" or file_path.endswith(".pdf"):
        if pypdf is not None:
            try:
                reader = pypdf.PdfReader(file_path)
                for i, page in enumerate(reader.pages):
                    text = page.extract_text() or ""
                    pages.append({"page": i + 1, "text": text})
            except Exception as e:
                print(f"Error parsing PDF with pypdf: {e}")
        else:
            print("Warning: pypdf is not installed. Unable to parse PDF. Please install it using 'pip install pypdf'")
            pages.append({"page": 1, "text": f"[PDF Document: PyPDF is not installed on the server to parse {os.path.basename(file_path)}]"})
            
    elif material_type == "docx" or file_path.endswith(".docx"):
        if docx is not None:
            try:
                doc = docx.Document(file_path)
                full_text = []
                for para in doc.paragraphs:
                    full_text.append(para.text)
                text = "\n".join(full_text)
                pages.append({"page": 1, "text": text})
            except Exception as e:
                print(f"Error parsing DOCX with python-docx: {e}")
        else:
            print("Warning: python-docx is not installed. Please install it using 'pip install python-docx'")
            pages.append({"page": 1, "text": f"[DOCX Document: python-docx is not installed on the server to parse {os.path.basename(file_path)}]"})
            
    elif material_type == "pptx" or file_path.endswith(".pptx"):
        if pptx is not None:
            try:
                prs = pptx.Presentation(file_path)
                for i, slide in enumerate(prs.slides):
                    slide_text = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            slide_text.append(shape.text)
                    text = "\n".join(slide_text)
                    pages.append({"page": i + 1, "text": text})
            except Exception as e:
                print(f"Error parsing PPTX with python-pptx: {e}")
        else:
            print("Warning: python-pptx is not installed. Please install it using 'pip install python-pptx'")
            pages.append({"page": 1, "text": f"[PPTX Document: python-pptx is not installed on the server to parse {os.path.basename(file_path)}]"})
            
    return pages

async def process_and_index_material(material_id: str, course_id: str, file_path: str, material_type: str):
    db = get_database()
    try:
        await db.materials.update_one(
            {"_id": ObjectId(material_id)},
            {"$set": {"vector_status": "processing"}}
        )

        pages = extract_text_from_file(file_path, material_type)
        if not pages:
            raise ValueError("No text could be extracted from the document.")

        all_chunks = []
        all_ids = []
        all_metadatas = []
        
        material = await db.materials.find_one({"_id": ObjectId(material_id)})
        material_title = material.get("title", "Unknown Document") if material else "Unknown Document"
        subject_id = str(material.get("subject_id", "")) if material else ""
        topic_id = str(material.get("topic_id", "")) if material else ""
        
        chunk_count = 0
        full_text_for_summary = ""
        
        for p_info in pages:
            page_num = p_info["page"]
            page_text = p_info["text"]
            full_text_for_summary += page_text + "\n"
            
            chunks = chunk_text(page_text, chunk_size=500, overlap=100)
            for chunk in chunks:
                if not chunk.strip():
                    continue
                chunk_count += 1
                chunk_id = f"{material_id}_p{page_num}_c{chunk_count}"
                
                all_chunks.append(chunk)
                all_ids.append(chunk_id)
                all_metadatas.append({
                    "material_id": str(material_id),
                    "course_id": str(course_id),
                    "subject_id": subject_id,
                    "topic_id": topic_id,
                    "material_title": material_title,
                    "page_number": page_num,
                    "text": chunk
                })

        if all_chunks:
            model = get_embedding_model()
            embeddings = model.encode(all_chunks).tolist() if model else []
            collection = chroma_client.get_or_create_collection("course_materials")
            if embeddings:
                collection.upsert(
                    ids=all_ids,
                    documents=all_chunks,
                    embeddings=embeddings,
                    metadatas=all_metadatas
                )

        summary = "No summary available."
        if full_text_for_summary.strip():
            summary_prompt = f"Write a short, professional, 2-3 sentence educational summary of the following syllabus material:\n\n{full_text_for_summary[:3000]}"
            system_instruction = "You are a professional teaching assistant. You write concise summaries of lecture materials."
            try:
                res = await LLMService.generate_text(summary_prompt, system_instruction)
                if res and not res.startswith("[") and "API Error" not in res and "Connection Error" not in res:
                    summary = res.strip()
                else:
                    summary = ""
            except Exception as se:
                print(f"Error generating summary: {se}")
                summary = ""

        await db.materials.update_one(
            {"_id": ObjectId(material_id)},
            {"$set": {
                "vector_status": "processed",
                "summary": summary
            }}
        )
        print(f"Successfully processed and indexed material {material_id} ({material_title})")
        
    except Exception as e:
        print(f"Error indexing material {material_id}: {e}")
        await db.materials.update_one(
            {"_id": ObjectId(material_id)},
            {"$set": {
                "vector_status": "failed",
                "summary": f"Failed to index: {str(e)}"
            }}
        )

async def query_rag_doubt(
    student_id: str, 
    course_id: str, 
    query: str,
    subject_name: Optional[str] = None,
    topic_name: Optional[str] = None,
    subtopic: Optional[str] = None
) -> dict:
    db = get_database()
    model = get_embedding_model()
    query_embedding = model.encode(query).tolist() if model else []
    collection = chroma_client.get_or_create_collection("course_materials")
    
    # Build learning context hierarchy string
    context_parts = []
    if subject_name:
        context_parts.append(subject_name)
    if topic_name:
        context_parts.append(topic_name)
    if subtopic:
        context_parts.append(subtopic)
    context_path = " → ".join(context_parts) if context_parts else ""

    # Dynamically determine if the ID is a subject_id, topic_id, or course_id
    where_filter = {}
    if ObjectId.is_valid(course_id):
        obj_id = ObjectId(course_id)
        if await db.subjects.find_one({"_id": obj_id}):
            where_filter = {"subject_id": str(course_id)}
        elif await db.topics.find_one({"_id": obj_id}):
            where_filter = {"topic_id": str(course_id)}
        else:
            where_filter = {"course_id": str(course_id)}
    else:
        where_filter = {"course_id": str(course_id)}
        
    search_results = None
    if query_embedding:
        try:
            search_results = collection.query(
                query_embeddings=[query_embedding],
                n_results=5,
                where=where_filter
            )
            # If no results found with specific filter, query collection globally
            if not search_results or not search_results.get("metadatas") or not search_results["metadatas"][0]:
                search_results = collection.query(
                    query_embeddings=[query_embedding],
                    n_results=5
                )
        except Exception as qe:
            print(f"ChromaDB query error ({qe}). Querying globally...")
            try:
                search_results = collection.query(
                    query_embeddings=[query_embedding],
                    n_results=5
                )
            except Exception:
                search_results = None
    
    retrieved_chunks = []
    if search_results and "metadatas" in search_results and search_results["metadatas"]:
        metadatas = search_results["metadatas"][0]
        documents = search_results["documents"][0]
        
        for meta, doc in zip(metadatas, documents):
            retrieved_chunks.append({
                "material_id": meta.get("material_id"),
                "material_title": meta.get("material_title", "Unknown Material"),
                "page_number": meta.get("page_number", 1),
                "text": meta.get("text", doc)
            })

    context_restriction_instruction = ""
    if context_path:
        context_restriction_instruction = f"""
IMPORTANT LEARNING CONTEXT RESTRICTION:
The student has explicitly selected their learning context as:
{context_path}

You must evaluate whether the student's question is relevant to this learning context ({context_path}).
If the question is unrelated (for example, asking about Computer Networks routing, Operating Systems paging, or cooking when the selected context is Java → Java Basics → Variables), you MUST NOT answer the question.
Instead, your response MUST be EXACTLY:
"Your current doubt-solving context is {context_path}. Please ask a question related to this topic or change the selected subject/topic."
Only answer the question if it is genuinely relevant to {context_path}.
"""

    if retrieved_chunks:
        context_str = ""
        for i, chunk in enumerate(retrieved_chunks):
            context_str += f"Source {i+1}: Document '{chunk['material_title']}', Page {chunk['page_number']}\nContent: {chunk['text']}\n\n"
            
        system_instruction = f"You are a professional educational Virtual AI Tutor. Answer the student's question based on educational concepts and provided course reference materials.{context_restriction_instruction}"
        prompt = f"""
Student Question: "{query}"

Here are the relevant sections from the course materials:
{context_str}

Please write a clear, helpful explanation.
"""
    else:
        system_instruction = f"You are a professional educational Virtual AI Tutor. Answer the student's question using clear, structured, academic explanations.{context_restriction_instruction}"
        prompt = f"Student Question: \"{query}\"\n\nExplain the concept clearly."

    response = await LLMService.generate_text(prompt, system_instruction)
    
    is_out_of_context = False
    if context_path and (
        "Your current doubt-solving context is" in response 
        or "Please ask a question related to this topic" in response
    ):
        is_out_of_context = True
        response = f"Your current doubt-solving context is {context_path}. Please ask a question related to this topic or change the selected subject/topic."
    
    citations = []
    for chunk in retrieved_chunks:
        citations.append({
            "material_id": ObjectId(chunk["material_id"]) if ObjectId.is_valid(chunk["material_id"]) else chunk["material_id"],
            "page_number": chunk["page_number"],
            "snippet": chunk["text"]
        })
        
    log_dict = {
        "student_id": ObjectId(student_id) if ObjectId.is_valid(student_id) else student_id,
        "query": query,
        "response": response,
        "citations": citations,
        "timestamp": datetime.now(timezone.utc)
    }
    # Log the filter correctly
    if "subject_id" in where_filter:
        log_dict["subject_id"] = ObjectId(course_id)
    elif "topic_id" in where_filter:
        log_dict["topic_id"] = ObjectId(course_id)
    else:
        log_dict["course_id"] = ObjectId(course_id) if ObjectId.is_valid(course_id) else course_id
    
    try:
        await db.rag_queries.insert_one(log_dict)
    except Exception as le:
        print(f"Error logging RAG query: {le}")

    if is_out_of_context:
        return {
            "response": response,
            "citations": [],
            "youtube_resources": []
        }

    # Generate real YouTube video tutorial link
    from urllib.parse import quote_plus
    search_term = query.strip()
    yt_query = quote_plus(f"{search_term} lecture tutorial")
    yt_search_url = f"https://www.youtube.com/results?search_query={yt_query}"
    
    youtube_resources = [
        {
            "title": f"Watch '{search_term[:45]}' Video Tutorial",
            "url": yt_search_url,
            "description": "Click to open relevant video lessons on YouTube in a new tab.",
            "channel": "YouTube Education"
        }
    ]

    return {
        "response": response,
        "citations": [
            {
                "material_id": str(chunk["material_id"]),
                "material_title": chunk["material_title"],
                "page_number": chunk["page_number"],
                "snippet": chunk["text"]
            } for chunk in retrieved_chunks
        ],
        "youtube_resources": youtube_resources
    }
